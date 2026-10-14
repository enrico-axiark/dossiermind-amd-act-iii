"""Portable deterministic extraction, without AI, network or temporary originals."""
import sys,json,csv,io,re,zipfile,hashlib
from decimal import Decimal,InvalidOperation
from xml.etree import ElementTree as ET
from email import policy
from email.parser import BytesParser

VERSION='act-iii-parser-0.1.2'
class Failure(Exception): pass
def extract(data,mime,name,depth=0):
    fragments=[];warnings=[]
    def fragment(text,locator,metadata=None):
        if text.strip(): fragments.append({'text':text,'locator':locator,'metadata':metadata or {}})
    if mime=='application/pdf' or name.lower().endswith('.pdf'):
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data),strict=True)
        if reader.is_encrypted: raise Failure('encrypted_pdf')
        if len(reader.pages)>50: raise Failure('page_limit_exceeded')
        for i,page in enumerate(reader.pages,1):
            text=page.extract_text() or ''
            if len(text)>100000:raise Failure('page_text_limit_exceeded')
            fragment(text,{'page':i})
            if not text.strip():warnings.append({'code':'needs_vision','page':i})
    elif mime in ('text/csv','application/csv') or name.lower().endswith('.csv'):
        text=data.decode('utf-8-sig',errors='strict');sample=text[:8192]
        try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
        except csv.Error:raise Failure('csv_dialect_unclear')
        rows=list(csv.reader(io.StringIO(text,newline=''),dialect))
        if not rows:raise Failure('empty_table')
        if len(rows)-1>10000:raise Failure('row_limit_exceeded')
        columns=rows[0]
        if not columns or len(set(columns))!=len(columns):raise Failure('ambiguous_columns')
        # Locator is the physical starting line, including quoted multiline cells.
        reader=csv.reader(io.StringIO(text,newline=''),dialect);next(reader);physical=reader.line_num+1
        for values in reader:
            if len(values)!=len(columns):raise Failure('inconsistent_columns')
            fields=dict(zip(columns,values));fragment(json.dumps(fields,ensure_ascii=False),{'row':physical},{'columns':fields,'delimiter':dialect.delimiter})
            for value in values:
                if re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}',value):warnings.append({'code':'ambiguous_date','row':physical})
                if value.startswith(('=','+','@')):warnings.append({'code':'formula_text','row':physical})
            physical=reader.line_num+1
    elif mime in ('application/xml','text/xml') or name.lower().endswith('.xml'):
        # Reject non-UTF8 to prevent bypass of DTD/entity checks via alternate encodings.
        text=data.decode('utf-8-sig',errors='strict')
        if '\0' in text:raise Failure('unsupported_xml_encoding')
        encoding=re.search(r'<\?xml[^>]*encoding=["\x27]([^"\x27]+)',text,re.I)
        if encoding and encoding[1].lower() not in ('utf-8','utf8','us-ascii'):raise Failure('unsupported_xml_encoding')
        if re.search(r'<!\s*(?:DOCTYPE|ENTITY)',text,re.I):raise Failure('unsafe_xml')
        root=ET.fromstring(text)
        count=0
        def visit(node,path,depth):
            nonlocal count
            count+=1
            if count>10000 or depth>64:raise Failure('xml_structure_limit')
            local=node.tag.split('}')[-1]
            if (node.text or '').strip():fragment(node.text.strip(),{'xml_path':path},{'tag':local})
            counts={}
            for child in node:
                tag=child.tag.split('}')[-1];counts[tag]=counts.get(tag,0)+1
                visit(child,path+'/'+tag+f'[{counts[tag]}]',depth+1)
        visit(root,'/'+root.tag.split('}')[-1],0)
    elif name.lower().endswith('.xlsx') or 'spreadsheetml' in mime:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos=archive.infolist()
            if len(infos)>10000 or sum(x.file_size for x in infos)>100*1024*1024:raise Failure('archive_limit_exceeded')
            for info in infos:
                if info.filename.startswith(('/', '\\')) or '..' in info.filename.replace('\\','/').split('/') or ':' in info.filename:raise Failure('unsafe_archive')
                if info.filename.lower().endswith('vbaproject.bin'):raise Failure('macros_not_allowed')
                if info.filename.lower().endswith('.xml'):
                    xml=archive.read(info).decode('utf-8-sig',errors='strict')
                    if '\0' in xml or re.search(r'<!\s*(?:DOCTYPE|ENTITY)',xml,re.I):raise Failure('unsafe_xml')
        from openpyxl import load_workbook
        formulas=load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
        cached=load_workbook(io.BytesIO(data),read_only=True,data_only=True,keep_links=False)
        try:
            if len(formulas.sheetnames)>20:raise Failure('sheet_limit_exceeded')
            for sheet in formulas:
                if (sheet.max_row or 0)>10000 or (sheet.max_column or 0)>200:raise Failure('sheet_size_limit_exceeded')
                # Read actual rows, not only producer-supplied dimensions. Write-only
                # and adversarial sheets can omit or understate that metadata.
                sheet.reset_dimensions()
                for row_number,row in enumerate(sheet.iter_rows(),1):
                    if row_number>10000 or len(row)>200:raise Failure('sheet_size_limit_exceeded')
                    for cell in row:
                        if cell.value is None:continue
                        value=cell.value;meta={}
                        if cell.data_type=='f':
                            meta={'formula':str(value),'cached_value':cached[sheet.title][cell.coordinate].value}
                            value=meta['cached_value']
                            warnings.append({'code':'formula_not_recalculated','sheet':sheet.title,'cell':cell.coordinate})
                            if value is None:warnings.append({'code':'formula_cache_missing','sheet':sheet.title,'cell':cell.coordinate})
                        fragment(str(value) if value is not None else str(meta['formula']),{'sheet':sheet.title,'cell':cell.coordinate},meta)
        finally:formulas.close();cached.close()
    elif mime=='message/rfc822' or name.lower().endswith('.eml'):
        message=BytesParser(policy=policy.default).parsebytes(data)
        if message.defects:warnings.append({'code':'email_structure_warning'})
        fragment('\n'.join(str(message.get(x,'')) for x in ['Subject','From','To','Date']),{'part':0},{'kind':'headers'})
        attachments=0
        for index,part in enumerate(message.walk()):
            if part.get_content_type()=='message/rfc822':raise Failure('nested_email_not_supported')
            if part.is_multipart():continue
            filename=part.get_filename()
            if filename:
                attachments+=1
                if attachments>20:raise Failure('attachment_limit_exceeded')
                if '/' in filename or '\\' in filename or '\0' in filename:
                    warnings.append({'code':'unsafe_attachment_name','part':index});continue
                if depth>=1:
                    warnings.append({'code':'attachment_depth_limit','part':index});continue
                payload=part.get_payload(decode=True) or b''
                if len(payload)>20*1024*1024:
                    warnings.append({'code':'attachment_size_limit','part':index});continue
                try:
                    attachment=extract(payload,part.get_content_type(),filename,depth+1)
                    for item in attachment['fragments']:
                        fragment(item['text'],{'part':index,'attachment':filename,**item['locator']},
                                 {**item['metadata'],'attachment_hash':hashlib.sha256(payload).hexdigest()})
                    warnings.append({'code':'attachment_extracted','part':index,'filename':filename,
                                     'status':attachment['status'],'fragment_count':len(attachment['fragments'])})
                    for warning in attachment['warnings']:warnings.append({**warning,'part':index,'attachment':filename})
                except Exception as error:
                    warnings.append({'code':'attachment_rejected','part':index,'filename':filename,
                                     'reason':str(error) if isinstance(error,Failure) else 'unreadable_document'})
                continue
            if part.get_content_type()=='text/plain':
                fragment(part.get_content(),{'part':index},{'kind':'body'})
            elif part.get_content_type()=='text/html':
                warnings.append({'code':'remote_resources_ignored','part':index})
                fragment(re.sub('<[^>]*>',' ',part.get_content()),{'part':index},{'kind':'html_text'})
    elif mime.startswith('image/'):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS=25000000
        image=Image.open(io.BytesIO(data));image.verify()
        warnings.append({'code':'needs_vision'})
    else:raise Failure('unsupported_format')
    if len(fragments)>30000:raise Failure('fragment_limit_exceeded')
    return {'parserVersion':VERSION,'fragments':fragments,'warnings':warnings,
            'status':'needs_vision' if any(x['code']=='needs_vision' for x in warnings) else 'extracted'}

try:
    data=sys.stdin.buffer.read(20*1024*1024+1)
    if len(data)>20*1024*1024:raise Failure('upload_too_large')
    result=extract(data,sys.argv[1],sys.argv[2])
    output=json.dumps(result,ensure_ascii=False,default=str)
    if len(output.encode())>8*1024*1024:raise Failure('extraction_output_limit')
    print(output)
except Exception as error:
    code=str(error) if isinstance(error,Failure) else 'parser_dependency_missing' if isinstance(error,ImportError) else 'unreadable_document'
    print(json.dumps({'error':code}));sys.exit(2)
